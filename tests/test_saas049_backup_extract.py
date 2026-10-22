"""SAAS-049 — look inside an old backup and export, without restoring."""

from __future__ import annotations

import io
import re

import pytest

from isolation_harness import host
from test_saas048_school_restore import _state

ACCOUNT_PASSWORD = "test-password-12345"     # the two_schools fixture's account password
PASS = "a long enough passphrase"


@pytest.fixture
def xapp(app, monkeypatch, tmp_path):
    monkeypatch.setitem(app.config, "TENANT_RESOLUTION_MODE", "enforce")
    monkeypatch.setitem(app.config, "WTF_CSRF_ENABLED", False)
    monkeypatch.setitem(app.config, "UPLOAD_ROOT", str(tmp_path / "uploads"))
    monkeypatch.setitem(app.config, "SCHOOL_BACKUP_DIR", str(tmp_path / "backups"))
    return app


def _client(app, login_as, bundle, role="admin"):
    c = app.test_client()
    login_as(c, bundle["users"][role], base_url=host(bundle))
    return c


def _saved(app, bundle):
    from app.extensions import db
    from app.models import SchoolBackup
    from app.services import school_backup as sb
    with app.app_context():
        res = sb.export_school(bundle["school"].id)
        b = db.session.get(SchoolBackup, bundle["backup"])
        b.file_name, b.size_bytes = res.path.name, res.bytes
        db.session.commit()
        db.session.remove()
    return bundle["backup"], res.path


def _open(c, bundle, bid):
    r = c.post("/admin/backups/browse", base_url=host(bundle), data={"backup_id": bid})
    assert r.status_code == 302, r.get_data(as_text=True)[:300]
    return r.headers["Location"].split("/admin/backups/browse/", 1)[1]


@pytest.mark.parametrize("role", ["teacher", "student", "parent"])
def test_only_the_owner(xapp, world, login_as, role):
    A = world["A"]
    c = _client(xapp, login_as, A, role)
    assert c.get("/admin/backups/browse", base_url=host(A)).status_code == 403
    assert c.post("/admin/backups/browse", base_url=host(A), data={"backup_id": A["backup"]}).status_code == 403


def test_browse_search_and_export_leave_live_data_untouched(xapp, world, login_as):
    from openpyxl import load_workbook
    A = world["A"]
    bid, _ = _saved(xapp, A)
    before = _state(xapp, A)
    c = _client(xapp, login_as, A)
    token = _open(c, A, bid)
    page = c.get(f"/admin/backups/browse/{token}", base_url=host(A)).get_data(as_text=True)
    assert "الطلاب" in page and "الأساسيات الأكاديمية" in page
    found = c.get(f"/admin/backups/browse/{token}?q=Student-alpha", base_url=host(A)).get_data(as_text=True)
    assert "PC-alpha-01" in found
    sid = A["student"].id
    t = c.get(f"/admin/backups/browse/{token}/t/enrollments?student={sid}", base_url=host(A))
    assert t.status_code == 200
    r = c.post(f"/admin/backups/browse/{token}/t/enrollments/export", base_url=host(A),
               data={"student": sid, "format": "xlsx", "account_password": ACCOUNT_PASSWORD})
    assert r.status_code == 200 and r.headers["Cache-Control"] == "no-store"
    ws = load_workbook(io.BytesIO(r.get_data())).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "id" and len(rows) == 2           # header + A's one enrollment
    assert _state(xapp, A) == before


def test_wrong_password_refused_and_audited(xapp, world, login_as):
    from app.extensions import db
    from app.models import AuditLog
    A = world["A"]
    bid, _ = _saved(xapp, A)
    c = _client(xapp, login_as, A)
    token = _open(c, A, bid)
    r = c.post(f"/admin/backups/browse/{token}/t/students/export", base_url=host(A), follow_redirects=True,
               data={"format": "csv", "account_password": "wrong"})
    assert "غير صحيحة" in r.get_data(as_text=True)
    with xapp.app_context():
        acts = {a.action for a in AuditLog.query.filter_by(school_id=A["school"].id)}
        db.session.remove()
    assert {"backup_browse", "extract_denied"} <= acts


def test_b_package_refused_and_its_upload_removed(xapp, world, login_as, tmp_path):
    from app.services import school_backup as sb
    A, B = world["A"], world["B"]
    _, path_b = _saved(xapp, B)
    enc = tmp_path / "b.mbk"
    sb.encrypt_file(path_b, enc, PASS)
    c = _client(xapp, login_as, A)
    r = c.post("/admin/backups/browse", base_url=host(A), follow_redirects=True,
               data={"package": (io.BytesIO(enc.read_bytes()), "b.mbk"), "passphrase": PASS},
               content_type="multipart/form-data")
    assert "مرفوضة" in r.get_data(as_text=True)
    with xapp.app_context():
        left = list((sb.backups_dir(A["school"].id) / "browse").glob("*"))
    assert left == []


def test_token_is_bound_to_the_user_and_school(xapp, world, login_as):
    from app.extensions import db
    from app.models import User
    A, B = world["A"], world["B"]
    bid, _ = _saved(xapp, A)
    token = _open(_client(xapp, login_as, A), A, bid)
    with xapp.app_context():
        other = User(school_id=A["school"].id, role_id=A["roles"]["admin"].id, username="admin3",
                     full_name="Admin3", is_active=True)
        other.password_hash = db.session.get(User, A["users"]["admin"].id).password_hash
        db.session.add(other)
        db.session.commit()
        db.session.refresh(other)
        db.session.expunge(other)
        db.session.remove()
    c2 = xapp.test_client()
    login_as(c2, other, base_url=host(A))
    assert c2.get(f"/admin/backups/browse/{token}", base_url=host(A)).status_code == 404
    cb = _client(xapp, login_as, B)
    assert cb.get(f"/admin/backups/browse/{token}", base_url=host(B)).status_code == 404


def test_uploaded_copies_are_swept(xapp, world, login_as, monkeypatch):
    from app.services import backup_extract as bx
    from app.services import school_backup as sb
    A = world["A"]
    _, path_a = _saved(xapp, A)
    c = _client(xapp, login_as, A)
    r = c.post("/admin/backups/browse", base_url=host(A),
               data={"package": (io.BytesIO(path_a.read_bytes()), "a.tar.gz")},
               content_type="multipart/form-data")
    assert r.status_code == 302
    with xapp.app_context():
        d = sb.backups_dir(A["school"].id) / "browse"
        assert len(list(d.glob("*.tar.gz"))) == 1
        monkeypatch.setattr(bx, "BROWSE_TTL", -1)
        assert bx.sweep(A["school"].id) == 1
        assert list(d.glob("*")) == []


def test_export_escapes_formulas_and_respects_the_row_cap(xapp, world, login_as, monkeypatch):
    from app.extensions import db
    from app.models import Student
    from app.services import backup_extract as bx
    A = world["A"]
    with xapp.app_context():
        db.session.add(Student(school_id=A["school"].id, permanent_code="PC-alpha-77", full_name="=HYPERLINK(1)"))
        db.session.commit()
        db.session.remove()
    bid, _ = _saved(xapp, A)
    c = _client(xapp, login_as, A)
    token = _open(c, A, bid)
    r = c.post(f"/admin/backups/browse/{token}/t/students/export", base_url=host(A),
               data={"format": "csv", "account_password": ACCOUNT_PASSWORD})
    text = r.get_data(as_text=True)
    assert "'=HYPERLINK(1)" in text and not re.search(r"(^|,)=HYPERLINK", text, re.M)
    monkeypatch.setattr(bx, "EXPORT_MAX_ROWS", 1)
    r = c.post(f"/admin/backups/browse/{token}/t/students/export", base_url=host(A),
               data={"format": "csv", "account_password": ACCOUNT_PASSWORD})
    assert len(r.get_data(as_text=True).strip().splitlines()) == 2      # header + 1
